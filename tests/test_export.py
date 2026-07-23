"""Document export tests (docs/13): SVG snapshot + determinism, DXF round-trip, XLSX content,
zip bundle completeness. No network; everything is deterministic."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import ezdxf
from openpyxl import load_workbook

from electricopilot.export import (
    build_boq,
    build_bundle,
    build_cable_journal,
    build_sld,
    sld_sheets_svg,
)
from electricopilot.export.dxf import render_dxf
from electricopilot.export.primitives import Circle, Drawing, Line, Polyline, Rect, Text
from electricopilot.export.svg import render_svg
from electricopilot.project import build_project_report


def _circuit(cid, ref, desc, P, ph=1, U=230, pf=1.0, purpose="power", method="C",
             material="Cu", insulation="PVC", L=20, dev="MCB", curve="C", iscc=1500,
             phase="L1", rcd=False):
    return {
        "id": cid, "ref": ref,
        "request": {
            "load": {"description": desc, "power_w": P, "voltage_v": U, "phases": ph,
                     "power_factor": pf, "purpose": purpose},
            "installation": {"method": method, "material": material, "insulation": insulation,
                             "ambient_temp_c": 30, "grouping_circuits": 1, "length_m": L},
            "protection": {"device_class": dev, "prospective_fault_current_a": iscc,
                           "trip_curve_type": curve},
        },
        "meta": {"phase": phase, "rcd": ({"present": True, "type": "RCBO", "ma": 30}
                                         if rcd else {"present": False})},
    }


def _board(n_circuits=3):
    circuits = [
        _circuit("c1", "L1", "Розетки кухни", 3680, U=230, pf=0.95, purpose="socket",
                 L=18, rcd=True),
        _circuit("c2", "L2", "Освещение", 1200, U=230, purpose="lighting", L=25, dev="MCB",
                 curve="B", iscc=800, phase="L2"),
        _circuit("c3", "L3", "Бойлер", 3000, U=230, L=20, rcd=True, phase="L3"),
    ][:n_circuits]
    return {
        "id": "project-export", "name": "Щит ВРУ-1 (пример)", "board_ref": "DB-1",
        "export_settings": {"cable_margin": 1.05},
        "supply": {"voltage_v": 400, "phases": 3, "ways_total": 12, "earthing": "TN-C-S"},
        "circuits": circuits,
    }


# --- SVG ---------------------------------------------------------------------------------
_SVG_SNAPSHOT = (
    '<svg xmlns="http://www.w3.org/2000/svg" width="100mm" height="50mm" '
    'viewBox="0 0 100 50" fill="none">\n'
    '<rect x="0" y="0" width="100" height="50" fill="#ffffff"/>\n'
    '<line x1="0" y1="0" x2="10" y2="10" stroke="#1f9d55" stroke-width="0.25"/>\n'
    '<rect x="5" y="5" width="20" height="10" stroke="#5b6b86" stroke-width="0.25" fill="none"/>\n'
    '<circle cx="30" cy="20" r="3" stroke="#0b1220" stroke-width="0.25" fill="none"/>\n'
    '<polygon points="0,0 5,0 2.5,5" stroke="#0b1220" stroke-width="0.25" fill="#0b1220"/>\n'
    '<text x="10" y="20" font-size="2.5" font-family="Helvetica, Arial, sans-serif" '
    'fill="#0b1220" text-anchor="middle" transform="rotate(-90 10 20)">QF &lt;1&gt;</text>\n'
    '</svg>'
)


def _tiny_drawing() -> Drawing:
    d = Drawing(100, 50)
    d.add(Line(0, 0, 10, 10, layer="WIRES", color="#1f9d55"))
    d.add(Rect(5, 5, 20, 10, layer="FRAME"))
    d.add(Circle(30, 20, 3, layer="SYMBOLS"))
    d.add(Polyline(((0, 0), (5, 0), (2.5, 5)), layer="SYMBOLS", closed=True))
    d.add(Text(10, 20, "QF <1>", height=2.5, anchor="middle", rotation=90))
    return d


def test_svg_golden_snapshot():
    """Deterministic SVG (escaped text, color override, rotation) matches a fixed snapshot."""
    assert render_svg(_tiny_drawing()) == _SVG_SNAPSHOT


def test_svg_is_deterministic():
    d = _tiny_drawing()
    assert render_svg(d) == render_svg(d)  # no random ids / timestamps


def test_sld_svg_has_key_content():
    board = _board()
    report = build_project_report(board)
    svgs = sld_sheets_svg(board, report)
    assert len(svgs) == 1
    svg = svgs[0]
    for token in ["L1", "L2", "L3", "Розетки кухни", "Ввод", "Однолинейная", "REVIEW"]:
        assert token in svg, token


def test_sld_splits_into_sheets_over_16():
    circuits = [_circuit(f"c{i}", f"C{i}", f"Цепь {i}", 2000) for i in range(17)]
    board = {**_board(0), "circuits": circuits}
    report = build_project_report(board)
    sheets = build_sld(board, report)
    assert len(sheets) == 2


# --- DXF ---------------------------------------------------------------------------------
def test_dxf_round_trip():
    board = _board()
    report = build_project_report(board)
    dxf_bytes = render_dxf(build_sld(board, report)[0])
    doc = ezdxf.read(io.StringIO(dxf_bytes.decode("utf-8")))
    msp = doc.modelspace()

    layers = {la.dxf.name for la in doc.layers}
    assert {"FRAME", "BUS", "WIRES", "SYMBOLS", "TEXT"} <= layers
    assert len(list(msp)) > 20  # frame + bus + per-circuit symbols/labels

    texts = [e.dxf.text for e in msp if e.dxftype() in ("TEXT", "MTEXT")]
    joined = " ".join(texts)
    for token in ["L1", "MCB", "мм²", "Розетки кухни"]:  # refs, device, section, description
        assert token in joined, token


def test_dxf_is_byte_deterministic():
    """DXF metadata is pinned without changing the R2010 geometry payload."""
    drawing = _tiny_drawing()
    assert render_dxf(drawing) == render_dxf(drawing)


# --- XLSX --------------------------------------------------------------------------------
def _read_sheet(xlsx_bytes: bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def test_cable_journal_xlsx_matches_rows():
    from electricopilot.export.xlsx import render_table_xlsx
    board = _board()
    report = build_project_report(board)
    table = build_cable_journal(board, report)
    grid = _read_sheet(render_table_xlsx(table))
    assert grid[0][0] == "Кабельный журнал"          # title
    assert grid[1][:3] == ["Обозначение", "Начало", "Конец"]  # header
    # one data row per circuit, with a with-margin length column
    data = grid[2:2 + len(report["rows"])]
    assert len(data) == len(report["rows"])
    for row, rep_row in zip(data, report["rows"]):
        assert row[0] == rep_row["ref"]
        length = float(rep_row["length_m"])
        assert abs(float(row[6]) - round(length * 1.05, 1)) < 0.05  # 5% margin applied


def test_cable_journal_margin_from_settings():
    board = _board()
    board["export_settings"] = {"cable_margin": 1.20}
    report = build_project_report(board)
    table = build_cable_journal(board, report)
    row0, rep0 = table.rows[0], report["rows"][0]
    assert abs(float(row0[6]) - round(float(rep0["length_m"]) * 1.20, 1)) < 0.05


def test_boq_groups_devices_and_cables():
    board = _board()
    report = build_project_report(board)
    table = build_boq(board, report)
    devices = [r for r in table.rows if r[1] == "Аппарат"]
    cables = [r for r in table.rows if r[1] == "Кабель"]
    assert devices and cables
    # every device count is шт, cable qty is м; total device count == circuit count
    assert all(r[3] == "шт" for r in devices)
    assert all(r[3] == "м" for r in cables)
    assert sum(r[4] for r in devices) == len(report["rows"])


# --- Bundle ------------------------------------------------------------------------------
def test_bundle_has_all_documents():
    board = _board()
    data = build_bundle(board)
    zf = zipfile.ZipFile(io.BytesIO(data))
    names = set(zf.namelist())
    assert {"report.md", "sld.svg", "sld.dxf",
            "cable_journal.xlsx", "boq.xlsx", "project.json"} <= names
    for name in names:
        assert zf.getinfo(name).file_size > 0, name
    # the embedded xlsx and dxf are themselves valid
    load_workbook(io.BytesIO(zf.read("cable_journal.xlsx")))
    ezdxf.read(io.StringIO(zf.read("sld.dxf").decode("utf-8")))


def test_pue_rk_trust_status_is_consistent_in_every_document_surface():
    board = _board(1)
    board["norm_pack"] = "pue-rk"
    report = build_project_report(board)
    expected = ["pue-rk", "v0.1.0", "public_standard", "NEEDS_REVIEW"]
    assert report["board"]["status"] == "NEEDS_REVIEW"
    assert report["norm_pack"]["verification_status"] == "NEEDS_REVIEW"
    assert "проверена инженером" not in report["provenance_note"]
    for token in [*expected, "UNSIGNED_ADVISORY"]:
        assert token in report["markdown"]

    drawing = build_sld(board, report)[0]
    svg = render_svg(drawing)
    dxf = ezdxf.read(io.StringIO(render_dxf(drawing).decode("utf-8")))
    dxf_text = " ".join(
        e.dxf.text for e in dxf.modelspace() if e.dxftype() in ("TEXT", "MTEXT")
    )
    for surface in (svg, dxf_text):
        for token in [*expected, "UNSIGNED_ADVISORY"]:
            assert token in surface
        assert report["disclaimer"] in surface

    for table in (build_cable_journal(board, report), build_boq(board, report)):
        notes = " ".join(table.notes)
        for token in [*expected, "UNSIGNED_ADVISORY"]:
            assert token in notes
        assert report["disclaimer"] in notes

    zf = zipfile.ZipFile(io.BytesIO(build_bundle(board)))
    for name in ("report.md", "sld.svg"):
        text = zf.read(name).decode("utf-8")
        for token in [*expected, "UNSIGNED_ADVISORY"]:
            assert token in text
    for name in ("cable_journal.xlsx", "boq.xlsx"):
        workbook = load_workbook(io.BytesIO(zf.read(name)))
        text = " ".join(
            str(cell)
            for row in workbook.active.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        )
        for token in [*expected, "UNSIGNED_ADVISORY"]:
            assert token in text
        assert report["disclaimer"] in text
    app_js = (Path(__file__).parents[1] / "web/app.js").read_text(encoding="utf-8")
    assert "rep.data_identity" in app_js and "rep.signoff_notice" in app_js


def test_bundle_multi_sheet_adds_extra_sld_files():
    circuits = [_circuit(f"c{i}", f"C{i}", f"Цепь {i}", 2000) for i in range(17)]
    board = {**_board(0), "circuits": circuits}
    zf = zipfile.ZipFile(io.BytesIO(build_bundle(board)))
    names = set(zf.namelist())
    assert "sld.svg" in names and "sld-2.svg" in names
    assert "sld.dxf" in names and "sld-2.dxf" in names


def test_bundle_is_byte_deterministic():
    """Fixed ZIP metadata and fixed DXF metadata make the complete bundle reproducible."""
    assert build_bundle(_board()) == build_bundle(_board())


def test_bundle_contains_typed_calculation_manifest_inputs():
    zf = zipfile.ZipFile(io.BytesIO(build_bundle(_board())))
    assert {"manifest.json", "calculation-input.json", "norm-pack.json"} <= set(zf.namelist())


# --- review-fix regressions -------------------------------------------------------------
def test_rcd_present_without_ma_defaults_to_30():
    """rcd.present=true but no `ma` must render '30мА', not 'NoneмА' (code-review finding)."""
    from electricopilot.export.boq import build_boq as _bq
    board = _board(1)
    board["circuits"][0]["meta"]["rcd"] = {"present": True, "type": "RCBO"}  # no ma
    report = build_project_report(board)
    dev = [r[2] for r in _bq(board, report).rows if r[1] == "Аппарат"][0]
    assert "None" not in dev and "30мА" in dev
    assert "None" not in sld_sheets_svg(board, report)[0]


def test_boq_distinguishes_poles():
    """A 1-phase and a 3-phase breaker of the same rating must be separate BoQ lines."""
    from electricopilot.export.boq import _device_name
    one = {"device_class": "MCB", "In_a": 16, "curve": "C", "phases": 1, "rcd": {"present": False}}
    three = {**one, "phases": 3}
    assert _device_name(one) != _device_name(three)
    assert "1P" in _device_name(one) and "3P" in _device_name(three)


def test_dxf_preserves_text_color():
    """Text colour must survive to DXF (status labels / disclaimer), not fall back to black."""
    d = Drawing(100, 50)
    d.add(Text(10, 10, "FAIL", height=3, color="#e5484d"))   # red status
    d.add(Text(10, 20, "footer", height=2, color="#556"))     # 3-digit hex
    doc = ezdxf.read(io.StringIO(render_dxf(d).decode("utf-8")))
    texts = {e.dxf.text: e for e in doc.modelspace() if e.dxftype() == "TEXT"}
    assert texts["FAIL"].rgb == (0xE5, 0x48, 0x4D)
    assert texts["footer"].rgb == (0x55, 0x55, 0x66)  # #556 → #555566


def test_xlsx_is_byte_deterministic():
    """openpyxl's wall-clock timestamps are pinned, so the same Table yields identical bytes,
    and no current wall-clock year leaks into the package."""
    from electricopilot.export.xlsx import render_table_xlsx
    from electricopilot.export.tables import Table
    t = Table("T", ["a", "b"], [[1, 2], [3, 4]], ["note"])
    a, b = render_table_xlsx(t), render_table_xlsx(t)
    assert a == b
    core = zipfile.ZipFile(io.BytesIO(a)).read("docProps/core.xml").decode("utf-8")
    assert "2020-01-01" in core and "2026" not in core


def test_boq_and_journal_cable_lengths_reconcile():
    """BoQ per-cable total equals the sum of the cable journal's per-row lengths (both round
    per row) — no 0.1 m drift between the two bundled documents."""
    board = _board()
    report = build_project_report(board)
    journal = build_cable_journal(board, report)
    boq = build_boq(board, report)
    # journal 'Длина с запасом, м' is column index 6; group by the section label (col 4)
    from collections import defaultdict
    by_section: dict[str, float] = defaultdict(float)
    for row in journal.rows:
        by_section[row[4]] += float(row[6])
    boq_cables = {r[2]: r[4] for r in boq.rows if r[1] == "Кабель"}
    # each BoQ cable line's metre total must equal the journal rows summed for that section
    for name, qty in boq_cables.items():
        section = name.split("Кабель ")[1].split(" мм²")[0] + " мм²"
        assert abs(round(by_section[section], 1) - qty) < 0.05, name
