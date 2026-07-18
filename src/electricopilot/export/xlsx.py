"""Render a Table to a single-sheet XLSX workbook via openpyxl (docs/13).

Layout: title row, bold header, data rows, then notes (disclaimer + provenance). Column widths
are sized to content. Byte-deterministic: openpyxl's wall-clock package timestamps (docProps
modified + every zip member date) are pinned to a fixed epoch by _normalize_xlsx.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .tables import Table

# Fixed epoch pinned into docProps so the workbook is byte-deterministic (openpyxl otherwise
# stamps datetime.now() into created/modified AND into every zip member's date, making the
# same Table differ run-to-run). _normalize_xlsx() below rewrites both.
_FIXED_DT = datetime(2020, 1, 1, tzinfo=timezone.utc)
_FIXED_ZIP_DATE = (2020, 1, 1, 0, 0, 0)
_MODIFIED_RE = re.compile(
    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>", re.DOTALL)
_FIXED_MODIFIED = (b'<dcterms:modified xsi:type="dcterms:W3CDTF">'
                   b"2020-01-01T00:00:00Z</dcterms:modified>")


def _normalize_xlsx(raw: bytes) -> bytes:
    """Rewrite the xlsx package with fixed member timestamps and a pinned core.xml modified date
    so the workbook is byte-identical across runs (openpyxl bakes wall-clock times into both)."""
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            data = src.read(name)
            if name == "docProps/core.xml":
                data = _MODIFIED_RE.sub(_FIXED_MODIFIED, data)
            info = zipfile.ZipInfo(filename=name, date_time=_FIXED_ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, data)
    return out.getvalue()

_HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
_TITLE_FONT = Font(bold=True, size=13)
_HEADER_FONT = Font(bold=True)
_NOTE_FONT = Font(italic=True, size=9, color="555555")


def render_table_xlsx(table: Table) -> bytes:
    wb = Workbook()
    wb.properties.created = _FIXED_DT
    wb.properties.modified = _FIXED_DT
    ws = wb.active
    ws.title = table.title[:31] or "Лист1"
    ncol = max(len(table.columns), 1)

    ws.append([table.title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.cell(row=1, column=1).font = _TITLE_FONT

    ws.append(table.columns)
    for c in range(1, ncol + 1):
        cell = ws.cell(row=2, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in table.rows:
        ws.append(list(row))

    ws.append([])
    for note in table.notes:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=note).font = _NOTE_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)

    # column widths from content
    for c in range(1, ncol + 1):
        header = table.columns[c - 1] if c - 1 < len(table.columns) else ""
        width = len(str(header))
        for row in table.rows:
            if c - 1 < len(row):
                width = max(width, len(str(row[c - 1])))
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 48)

    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return _normalize_xlsx(buf.getvalue())
