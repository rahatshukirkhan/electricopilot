"""Deterministic XLSX/CSV schedule import and R11 declared-vs-sized diff (docs/14 §8)."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Literal, Mapping, Protocol, TypeAlias
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import BaseModel

from .data.loader import DataPack
from .engine import size
from .exceptions import LlmConfigError, LlmError
from .llm.client import parse_json_lenient
from .models import DISCLAIMER
from .project_contract import Project, project_payload, validate_project

MAX_IMPORT_BYTES = 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 8 * MAX_IMPORT_BYTES
MAX_IMPORT_ROWS = 1000
MAX_IMPORT_COLUMNS = 64

CanonicalField = Literal[
    "ref",
    "description",
    "power_kw",
    "voltage_v",
    "phases",
    "power_factor",
    "length_m",
    "device_class",
    "declared_in_a",
    "trip_curve_type",
    "rcd_ma",
    "declared_section_mm2",
    "material",
    "insulation",
    "installation_method",
]
MappingMethod = Literal["heuristic", "llm", "manual", "unmapped"]
Scalar: TypeAlias = str | int | float | bool | None

CANONICAL_FIELDS: tuple[CanonicalField, ...] = (
    "ref", "description", "power_kw", "voltage_v", "phases", "power_factor",
    "length_m", "device_class", "declared_in_a", "trip_curve_type", "rcd_ma",
    "declared_section_mm2", "material", "insulation", "installation_method",
)


class ScheduleImportError(ValueError):
    def __init__(self, code: str, message: str, *, status_code: int = 422) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


class HeaderMappingClient(Protocol):
    def complete(
        self, *, model: str, system: str, user: str,
        json_schema: dict[str, Any] | None = None, max_output_tokens: int = 2048,
    ) -> str: ...


class MappingEntry(BaseModel):
    source_header: str
    field: CanonicalField | None = None
    method: MappingMethod


class ImportIssue(BaseModel):
    code: str
    message: str
    severity: Literal["warning", "error", "info"] = "warning"
    row: int | None = None
    field: str | None = None
    assumed: bool = False


class ImportDiffCheck(BaseModel):
    field: Literal["declared_section_mm2", "declared_in_a"]
    status: Literal["match", "violation", "not_checked"]
    severity: Literal["error", "warning"]
    observed: float | None
    required: float
    reason: str | None = None


class ImportDiffRow(BaseModel):
    rule_id: Literal["R11"] = "R11"
    circuit_id: str
    circuit_ref: str
    status: Literal["match", "violation", "not_checked"]
    checks: list[ImportDiffCheck]
    observed: dict[str, float | None]
    required: dict[str, float]
    source_trusted: bool
    data_provenance: dict[str, Any]


class ImportDiffReport(BaseModel):
    rule_id: Literal["R11"] = "R11"
    rows: list[ImportDiffRow]
    summary: dict[str, int]
    data_identity: str
    disclaimer: str = DISCLAIMER
    signoff_notice: str = (
        "UNSIGNED_ADVISORY — требуется проверка и подпись квалифицированного инженера."
    )


class ImportScheduleResult(BaseModel):
    ok: bool = True
    project_draft: dict[str, Any]
    headers: list[str]
    mapping: list[MappingEntry]
    canonical_fields: list[str]
    preview: list[dict[str, Scalar]]
    issues: list[ImportIssue]
    requires_confirmation: bool
    assumptions_confirmed: bool
    diff: ImportDiffReport


@dataclass(frozen=True)
class ParsedSchedule:
    source_format: Literal["csv", "xlsx"]
    source_sheet: str | None
    headers: tuple[str, ...]
    rows: tuple[dict[str, Scalar], ...]


def _normal_header(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value).casefold().replace("ё", "е")
    normalized = normalized.replace("²", "2").replace("φ", " phi ")
    return " ".join(re.sub(r"[^\w]+", " ", normalized).split())


_SYNONYMS: dict[CanonicalField, tuple[str, ...]] = {
    "ref": ("ref", "номер", "обозначение", "цепь", "circuit", "circuit ref", "no"),
    "description": ("описание", "наименование", "нагрузка", "потребитель", "description", "load"),
    "power_kw": ("p квт", "мощность квт", "мощность", "power kw", "kw", "квт"),
    "voltage_v": ("u в", "напряжение в", "напряжение", "voltage v", "voltage"),
    "phases": ("фазы", "фазность", "число фаз", "phases", "phase", "ph"),
    "power_factor": ("cos phi", "cosfi", "cos ф", "коэффициент мощности", "power factor", "pf"),
    "length_m": ("l м", "длина м", "длина", "length m", "length"),
    "device_class": ("аппарат", "тип аппарата", "device", "device class", "breaker type"),
    "declared_in_a": ("in а", "номинал а", "номинал автомата", "ток автомата", "breaker rating", "in a"),
    "trip_curve_type": ("кривая", "характеристика", "curve", "trip curve"),
    "rcd_ma": ("узо ма", "узо", "дифток", "rcd ma", "rcd", "idn"),
    "declared_section_mm2": ("сечение мм2", "сечение", "cable section mm2", "cable section", "section mm2"),
    "material": ("материал", "материал жилы", "material", "conductor"),
    "insulation": ("изоляция", "insulation"),
    "installation_method": ("способ прокладки", "метод прокладки", "метод", "installation method", "method"),
}
_NORMAL_SYNONYMS = {
    _normal_header(alias): field
    for field, aliases in _SYNONYMS.items()
    for alias in aliases
}


def heuristic_mapping(headers: tuple[str, ...]) -> list[MappingEntry]:
    used: set[CanonicalField] = set()
    entries: list[MappingEntry] = []
    for header in headers:
        field = _NORMAL_SYNONYMS.get(_normal_header(header))
        if field in used:
            field = None
        if field is not None:
            used.add(field)
        entries.append(MappingEntry(
            source_header=header,
            field=field,
            method="heuristic" if field is not None else "unmapped",
        ))
    return entries


def _validated_mapping(
    headers: tuple[str, ...], raw: Mapping[str, object], method: MappingMethod,
) -> list[MappingEntry]:
    unknown_headers = sorted(set(raw) - set(headers))
    if unknown_headers:
        raise ScheduleImportError(
            "unknown_mapping_header",
            f"Mapping содержит неизвестные заголовки: {', '.join(unknown_headers)}.",
        )
    used: dict[str, str] = {}
    entries: list[MappingEntry] = []
    for header in headers:
        value = raw.get(header)
        if value in (None, ""):
            field = None
        elif isinstance(value, str) and value in CANONICAL_FIELDS:
            field = value
        else:
            raise ScheduleImportError(
                "unknown_mapping_field", f"Недопустимое поле mapping для «{header}»: {value!r}.",
            )
        if field is not None and field in used:
            raise ScheduleImportError(
                "duplicate_mapping_field",
                f"Поле {field} назначено и «{used[field]}», и «{header}».",
            )
        if field is not None:
            used[field] = header
        entries.append(MappingEntry(source_header=header, field=field, method=method))
    return entries


_HEADER_MAPPING_SYSTEM = (
    "Сопоставь только названия колонок электрической ведомости с закрытым списком полей. "
    "Ты не получаешь значения строк и не должен создавать или вычислять их. Верни только JSON "
    "вида {\"mapping\": {\"исходный заголовок\": \"canonical_field\" или null}}."
)


def llm_header_mapping(
    headers: tuple[str, ...], client: HeaderMappingClient, *, model: str,
) -> list[MappingEntry]:
    raw = client.complete(
        model=model,
        system=_HEADER_MAPPING_SYSTEM,
        user=json.dumps(
            {"headers": headers, "allowed_fields": CANONICAL_FIELDS},
            ensure_ascii=False,
        ),
        json_schema={"type": "object"},
        max_output_tokens=1024,
    )
    parsed = parse_json_lenient(raw)
    if not isinstance(parsed, dict) or not isinstance(parsed.get("mapping"), dict):
        raise ValueError("LLM mapping must contain an object named mapping")
    return _validated_mapping(headers, parsed["mapping"], "llm")


def _merge_llm_mapping(
    heuristic: list[MappingEntry], llm: list[MappingEntry],
) -> list[MappingEntry]:
    llm_by_header = {entry.source_header: entry for entry in llm}
    used = {entry.field for entry in heuristic if entry.field is not None}
    merged: list[MappingEntry] = []
    for entry in heuristic:
        suggestion = llm_by_header[entry.source_header]
        if entry.field is None and suggestion.field is not None and suggestion.field not in used:
            merged.append(suggestion)
            used.add(suggestion.field)
        else:
            merged.append(entry)
    return merged


def _scalar(value: Any) -> Scalar:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def _unique_headers(values: list[Any]) -> tuple[str, ...]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value not in (None, "") else f"Column {index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise ScheduleImportError(
            "too_many_columns", f"Допустимо не более {MAX_IMPORT_COLUMNS} колонок.",
            status_code=413,
        )
    return tuple(headers)


def _rows_from_matrix(matrix: list[list[Any]]) -> tuple[tuple[str, ...], tuple[dict[str, Scalar], ...]]:
    nonempty = [row for row in matrix if any(value not in (None, "") for value in row)]
    if len(nonempty) < 2:
        raise ScheduleImportError("empty_schedule", "Файл не содержит заголовок и строки данных.")
    headers = _unique_headers(nonempty[0])
    data_rows = nonempty[1:]
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise ScheduleImportError(
            "too_many_rows", f"Допустимо не более {MAX_IMPORT_ROWS} строк.", status_code=413,
        )
    rows = tuple({
        header: _scalar(row[index] if index < len(row) else None)
        for index, header in enumerate(headers)
    } for row in data_rows)
    return headers, rows


def _parse_csv(data: bytes) -> ParsedSchedule:
    text: str | None = None
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ScheduleImportError("invalid_encoding", "CSV должен быть UTF-8 или Windows-1251.")
    sample = text[:8192]
    first_line = sample.splitlines()[0] if sample.splitlines() else ""
    if ";" in first_line or "\t" in first_line:
        delimiter = "\t" if first_line.count("\t") > first_line.count(";") else ";"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",").delimiter
        except csv.Error:
            delimiter = ","
    try:
        matrix = [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    except csv.Error as exc:
        raise ScheduleImportError("corrupt_csv", f"CSV повреждён: {exc}.") from None
    headers, rows = _rows_from_matrix(matrix)
    return ParsedSchedule(source_format="csv", source_sheet=None, headers=headers, rows=rows)


def _parse_xlsx(data: bytes) -> ParsedSchedule:
    stream = io.BytesIO(data)
    if not zipfile.is_zipfile(stream):
        raise ScheduleImportError("corrupt_xlsx", "XLSX не является корректным ZIP-контейнером.")
    stream.seek(0)
    try:
        with zipfile.ZipFile(stream) as archive:
            infos = archive.infolist()
            if len(infos) > 256 or sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ScheduleImportError(
                    "xlsx_expansion_limit", "Распакованный XLSX превышает безопасный лимит.",
                    status_code=413,
                )
        stream.seek(0)
        workbook = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = workbook.active
            sheet_title = sheet.title
            matrix: list[list[Any]] = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index > MAX_IMPORT_ROWS + 1:
                    raise ScheduleImportError(
                        "too_many_rows", f"Допустимо не более {MAX_IMPORT_ROWS} строк.",
                        status_code=413,
                    )
                matrix.append(list(row))
        finally:
            workbook.close()
        headers, rows = _rows_from_matrix(matrix)
        return ParsedSchedule(
            source_format="xlsx", source_sheet=sheet_title, headers=headers, rows=rows,
        )
    except ScheduleImportError:
        raise
    except Exception as exc:
        raise ScheduleImportError("corrupt_xlsx", f"Не удалось прочитать XLSX: {exc}.") from None


def parse_schedule(filename: str, data: bytes) -> ParsedSchedule:
    if len(data) > MAX_IMPORT_BYTES:
        raise ScheduleImportError(
            "file_too_large", "Файл превышает лимит 1 МиБ.", status_code=413,
        )
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        return _parse_csv(data)
    if suffix == ".xlsx":
        return _parse_xlsx(data)
    raise ScheduleImportError(
        "unsupported_file_type", "Поддерживаются только .xlsx и .csv.", status_code=415,
    )


def _number(value: Scalar) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value) if float(value) > 0 else None
    compact = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    match = re.fullmatch(r"([+]?(?:\d+(?:\.\d+)?|\.\d+))(?:[a-zа-я0-9²]+)?", compact.casefold())
    if not match:
        return None
    number = float(match.group(1))
    return number if number > 0 else None


def _enum(value: Scalar, aliases: Mapping[str, str], allowed: tuple[str, ...]) -> str | None:
    if value in (None, ""):
        return None
    normalized = _normal_header(str(value))
    resolved = aliases.get(normalized, str(value).strip())
    return resolved if resolved in allowed else None


def _mapped_values(row: dict[str, Scalar], mapping: list[MappingEntry]) -> dict[CanonicalField, Scalar]:
    return {
        entry.field: row.get(entry.source_header)
        for entry in mapping
        if entry.field is not None
    }


def _assumed(
    issues: list[ImportIssue], assumed_fields: list[str], *, row: int, field: str, message: str,
) -> None:
    assumed_fields.append(field)
    issues.append(ImportIssue(
        code="assumed_value", message=message, row=row, field=field, assumed=True,
    ))


def _mapping_method(entries: list[MappingEntry]) -> Literal["heuristic", "llm", "manual", "mixed"]:
    methods = {entry.method for entry in entries if entry.field is not None}
    if methods == {"manual"}:
        return "manual"
    if methods <= {"heuristic"}:
        return "heuristic"
    if methods <= {"llm"}:
        return "llm"
    return "mixed"


def build_import_project(
    parsed: ParsedSchedule,
    mapping: list[MappingEntry],
    *,
    filename: str,
    source_sha256: str,
    confirmed: bool,
    norm_pack: str | None,
) -> tuple[Project, list[ImportIssue]]:
    project_assumptions = [
        "norm_pack", "supply.voltage_v", "supply.phases", "supply.ways_total",
        "supply.earthing", "diversity.factors", "export_settings.cable_margin",
    ]
    issues: list[ImportIssue] = [ImportIssue(
        code="assumed_project_defaults",
        message=(
            "Параметры щита не пришли из таблицы и требуют проверки: "
            + ", ".join(project_assumptions) + "."
        ),
        assumed=True,
    )]
    circuits: list[dict[str, Any]] = []
    phase_names = ("L1", "L2", "L3")
    for index, raw_row in enumerate(parsed.rows):
        source_row = index + 2
        values = _mapped_values(raw_row, mapping)
        assumed_fields = [
            "request.load.purpose",
            "request.installation.ambient_temp_c",
            "request.installation.grouping_circuits",
            "request.protection.disconnection_time_s",
            "meta.phase",
        ]
        issues.append(ImportIssue(
            code="assumed_circuit_defaults",
            message=(
                "Не импортированы и требуют проверки: purpose, ambient_temp_c, "
                "grouping_circuits, disconnection_time_s и назначенная фаза."
            ),
            row=source_row,
            assumed=True,
        ))
        ref = str(values.get("ref") or f"C{index + 1}").strip()
        if not values.get("ref"):
            _assumed(issues, assumed_fields, row=source_row, field="ref", message="Ref создан по номеру строки.")
        description = str(values.get("description") or ref or f"Импортированная цепь {index + 1}").strip()
        if not values.get("description"):
            _assumed(issues, assumed_fields, row=source_row, field="description", message="Описание заменено на Ref.")

        power_kw = _number(values.get("power_kw"))
        if power_kw is None:
            power_kw = 2.0
            _assumed(issues, assumed_fields, row=source_row, field="power_kw", message="Мощность не распознана; применён UI-default 2 кВт.")
        voltage_v = _number(values.get("voltage_v"))
        if voltage_v is None:
            voltage_v = 230.0
            _assumed(issues, assumed_fields, row=source_row, field="voltage_v", message="Напряжение не распознано; применён UI-default 230 В.")
        phases_number = _number(values.get("phases"))
        phases = int(phases_number) if phases_number in (1.0, 3.0) else 1
        if phases_number not in (1.0, 3.0):
            _assumed(issues, assumed_fields, row=source_row, field="phases", message="Фазность не распознана; применён UI-default 1ф.")
        power_factor = _number(values.get("power_factor"))
        if power_factor is None or power_factor > 1:
            power_factor = 0.9
            _assumed(issues, assumed_fields, row=source_row, field="power_factor", message="cosφ не распознан; применён UI-default 0,9.")
        length_m = _number(values.get("length_m"))
        if length_m is None:
            length_m = 20.0
            _assumed(issues, assumed_fields, row=source_row, field="length_m", message="Длина не распознана; применён UI-default 20 м.")

        material = _enum(values.get("material"), {
            "медь": "Cu", "copper": "Cu", "cu": "Cu",
            "алюминий": "Al", "aluminium": "Al", "aluminum": "Al", "al": "Al",
        }, ("Cu", "Al"))
        if material is None:
            material = "Cu"
            _assumed(issues, assumed_fields, row=source_row, field="material", message="Материал не распознан; применён UI-default Cu.")
        insulation = _enum(values.get("insulation"), {
            "пвх": "PVC", "pvc": "PVC", "спэ": "XLPE", "xlpe": "XLPE",
        }, ("PVC", "XLPE"))
        if insulation is None:
            insulation = "PVC"
            _assumed(issues, assumed_fields, row=source_row, field="insulation", message="Изоляция не распознана; применён UI-default PVC.")
        method = _enum(values.get("installation_method"), {},
                       ("A1", "A2", "B1", "B2", "C", "D", "E", "F", "G"))
        if method is None:
            method = "C"
            _assumed(issues, assumed_fields, row=source_row, field="installation_method", message="Метод не распознан; применён UI-default C.")
        device = _enum(values.get("device_class"), {
            "автомат": "MCB", "автоматический выключатель": "MCB", "mcb": "MCB",
            "mccb": "MCCB", "предохранитель": "gG_fuse", "gg fuse": "gG_fuse",
        }, ("MCB", "MCCB", "gG_fuse"))
        if device is None:
            device = "MCB"
            _assumed(issues, assumed_fields, row=source_row, field="device_class", message="Аппарат не распознан; применён UI-default MCB.")
        curve = _enum(values.get("trip_curve_type"), {}, ("B", "C", "D"))
        if curve is None:
            curve = "C"
            _assumed(issues, assumed_fields, row=source_row, field="trip_curve_type", message="Кривая не распознана; применён UI-default C.")

        declared_section = _number(values.get("declared_section_mm2"))
        declared_in = _number(values.get("declared_in_a"))
        if declared_section is None:
            issues.append(ImportIssue(code="missing_declared_value", message="Нет заявленного сечения для R11.", row=source_row, field="declared_section_mm2"))
        if declared_in is None:
            issues.append(ImportIssue(code="missing_declared_value", message="Нет заявленного номинала для R11.", row=source_row, field="declared_in_a"))
        rcd_ma = _number(values.get("rcd_ma"))

        circuits.append({
            "id": f"ckt_{uuid4().hex[:12]}",
            "ref": ref,
            "sort_index": index,
            "request": {
                "load": {
                    "description": description,
                    "power_w": power_kw * 1000.0,
                    "voltage_v": voltage_v,
                    "phases": phases,
                    "power_factor": power_factor,
                    "purpose": "general",
                },
                "installation": {
                    "method": method,
                    "material": material,
                    "insulation": insulation,
                    "ambient_temp_c": 30.0,
                    "grouping_circuits": 1,
                    "length_m": length_m,
                },
                "protection": {
                    "device_class": device,
                    "prospective_fault_current_a": None,
                    "disconnection_time_s": 0.1,
                    "max_voltage_drop_pct": None,
                    "trip_curve_type": curve,
                },
            },
            "meta": {
                "phase": "L1L2L3" if phases == 3 else phase_names[index % 3],
                "rcd": ({"present": True, "type": "RCD", "ma": rcd_ma}
                        if rcd_ma is not None else {"present": False}),
                "diversity_category": "general",
                "import_declaration": {
                    "source_file": filename,
                    "source_sheet": parsed.source_sheet,
                    "source_row": source_row,
                    "raw_values": raw_row,
                    "declared_section_mm2": declared_section,
                    "declared_in_a": declared_in,
                    "assumed_fields": sorted(set(assumed_fields)),
                    "mapping_confirmed": confirmed,
                },
            },
            "result": None,
            "signoff": {"status": "UNSIGNED_ADVISORY"},
        })

    mapping_payload = {entry.source_header: entry.field for entry in mapping}
    project = validate_project({
        "schema_version": 2,
        "id": f"prj_{uuid4().hex[:12]}",
        "name": f"Импорт: {Path(filename).stem}",
        "board_ref": "IMPORTED",
        "location": "",
        "norm_pack": norm_pack,
        "supply": {
            "voltage_v": 400.0,
            "phases": 3,
            "ways_total": max(1, len(circuits)),
            "earthing": "TN-C-S",
            "method": "C",
            "material": "Cu",
            "insulation": "PVC",
            "ambient_temp_c": 30.0,
        },
        "diversity": {"factors": {}},
        "export_settings": {"cable_margin": 1.05},
        "circuits": circuits,
        "import_info": {
            "source_file": filename,
            "source_format": parsed.source_format,
            "source_sha256": source_sha256,
            "mapping": mapping_payload,
            "assumed_fields": project_assumptions,
            "assumptions_confirmed": confirmed,
            "mapping_method": _mapping_method(mapping),
        },
    })
    return project, issues


def _diff_check(
    *, field: Literal["declared_section_mm2", "declared_in_a"],
    observed: float | None, required: float, severity: Literal["error", "warning"],
    mismatch: bool, source_trusted: bool, sizing_failed: bool,
) -> ImportDiffCheck:
    if observed is None:
        return ImportDiffCheck(
            field=field, status="not_checked", severity=severity,
            observed=None, required=required, reason="missing_input",
        )
    if sizing_failed:
        return ImportDiffCheck(
            field=field, status="not_checked", severity=severity,
            observed=observed, required=required, reason="sizing_failed",
        )
    if not source_trusted:
        return ImportDiffCheck(
            field=field, status="not_checked", severity=severity,
            observed=observed, required=required, reason="untrusted_source",
        )
    return ImportDiffCheck(
        field=field,
        status="violation" if mismatch else "match",
        severity=severity,
        observed=observed,
        required=required,
    )


def build_import_diff(project: Project | Mapping[str, Any], pack: DataPack) -> ImportDiffReport:
    canonical = validate_project(project)
    rows: list[ImportDiffRow] = []
    for circuit in canonical.circuits:
        declaration = circuit.meta.import_declaration
        result = size(circuit.request, data_pack=pack)
        declared_section = declaration.declared_section_mm2 if declaration else None
        declared_in = declaration.declared_in_a if declaration else None
        required_section = result.selected_cable.cross_section_mm2
        required_in = result.selected_protection.In_a
        trusted = result.data_provenance.verification_status == "VERIFIED"
        sizing_failed = result.overall_status == "FAIL"
        checks = [
            _diff_check(
                field="declared_section_mm2", observed=declared_section,
                required=required_section, severity="error",
                mismatch=declared_section is not None and declared_section < required_section,
                source_trusted=trusted, sizing_failed=sizing_failed,
            ),
            _diff_check(
                field="declared_in_a", observed=declared_in,
                required=required_in, severity="warning",
                mismatch=declared_in is not None and declared_in != required_in,
                source_trusted=trusted, sizing_failed=sizing_failed,
            ),
        ]
        statuses = {check.status for check in checks}
        status: Literal["match", "violation", "not_checked"]
        status = "violation" if "violation" in statuses else (
            "not_checked" if "not_checked" in statuses else "match"
        )
        rows.append(ImportDiffRow(
            circuit_id=circuit.id,
            circuit_ref=circuit.ref,
            status=status,
            checks=checks,
            observed={
                "declared_section_mm2": declared_section,
                "declared_in_a": declared_in,
            },
            required={
                "calculated_section_mm2": required_section,
                "calculated_in_a": required_in,
            },
            source_trusted=trusted,
            data_provenance=result.data_provenance.model_dump(),
        ))
    summary = {name: sum(row.status == name for row in rows)
               for name in ("match", "violation", "not_checked")}
    return ImportDiffReport(
        rows=rows,
        summary=summary,
        data_identity=(
            f"Норм-пакет: {pack.meta.name} v{pack.meta.version}; "
            f"происхождение: {pack.meta.status}."
        ),
    )


def import_schedule(
    *,
    filename: str,
    data: bytes,
    pack: DataPack,
    manual_mapping: Mapping[str, object] | None = None,
    confirmed: bool = False,
    llm_client: HeaderMappingClient | None = None,
    llm_model: str | None = None,
) -> ImportScheduleResult:
    parsed = parse_schedule(filename, data)
    if manual_mapping is not None:
        mapping = _validated_mapping(parsed.headers, manual_mapping, "manual")
    else:
        mapping = heuristic_mapping(parsed.headers)
        if llm_client is not None and llm_model:
            try:
                mapping = _merge_llm_mapping(
                    mapping, llm_header_mapping(parsed.headers, llm_client, model=llm_model),
                )
            except (LlmConfigError, LlmError, ValueError, TypeError, json.JSONDecodeError):
                pass
    sha256 = hashlib.sha256(data).hexdigest()
    project, issues = build_import_project(
        parsed,
        mapping,
        filename=filename,
        source_sha256=sha256,
        confirmed=confirmed,
        norm_pack=pack.meta.name,
    )
    import_info = project.import_info
    assert import_info is not None
    requires_confirmation = bool(issues or import_info.assumed_fields)
    return ImportScheduleResult(
        project_draft=project_payload(project),
        headers=list(parsed.headers),
        mapping=mapping,
        canonical_fields=list(CANONICAL_FIELDS),
        preview=[dict(row) for row in parsed.rows[:5]],
        issues=issues,
        requires_confirmation=requires_confirmation,
        assumptions_confirmed=confirmed,
        diff=build_import_diff(project, pack),
    )
